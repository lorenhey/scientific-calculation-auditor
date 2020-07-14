import typer
from rich.console import Console
from rich.table import Table
from rich.panel import Panel
from pathlib import Path
import json

from .ir import AuditReport
from .auditors.python_auditor import PythonAuditor
from .auditors.notebook_auditor import NotebookAuditor
from .auditors.spreadsheet_auditor import SpreadsheetAuditor

app = typer.Typer(help="scientific-calculation-auditor: Audits scientific calculations.")
console = Console()

def print_report(report: AuditReport):
    if not report.findings:
        console.print(f"[bold green]No issues found in {report.target}[/bold green]")
        return
        
    console.print(f"[bold]Audit Report for {report.target}[/bold]\n")
    
    table = Table(show_header=True, header_style="bold magenta")
    table.add_column("Severity")
    table.add_column("Rule")
    table.add_column("Message")
    table.add_column("Evidence")
    
    for f in report.findings:
        sev_color = "red" if f.severity == "ERROR" or f.severity == "BLOCKING" else "yellow" if f.severity == "WARNING" else "blue"
        evidence_str = ""
        if f.evidence.line_number:
            evidence_str += f"Line {f.evidence.line_number} "
        if f.evidence.cell_reference:
            evidence_str += f"{f.evidence.cell_reference} "
        if f.evidence.context:
            evidence_str += f"\nContext: {f.evidence.context}"
            
        table.add_row(
            f"[{sev_color}]{f.severity.value}[/{sev_color}]",
            f.rule_id,
            f.message,
            evidence_str
        )
        
    console.print(table)
    
    errors = sum(1 for f in report.findings if f.severity in ("ERROR", "BLOCKING"))
    warnings = sum(1 for f in report.findings if f.severity == "WARNING")
    console.print(f"\n[bold red]{errors} errors[/bold red], [bold yellow]{warnings} warnings[/bold yellow]")

@app.command()
def audit(file_path: Path = typer.Argument(..., help="Path to file to audit"), fail_on_error: bool = False):
    """Audit a calculation artifact (Python script, Jupyter Notebook, or Excel Spreadsheet)."""
    if not file_path.exists():
        console.print(f"[red]File {file_path} not found.[/red]")
        raise typer.Exit(1)
        
    if file_path.suffix == ".py":
        auditor = PythonAuditor(file_path)
        report = auditor.audit()
    elif file_path.suffix == ".ipynb":
        auditor = NotebookAuditor(file_path)
        report = auditor.audit()
    elif file_path.suffix == ".xlsx":
        auditor = SpreadsheetAuditor(file_path)
        report = auditor.audit()
    else:
        console.print(f"[red]Unsupported file extension: {file_path.suffix}[/red]")
        raise typer.Exit(1)
        
    print_report(report)
    
    if fail_on_error and any(f.severity in ("ERROR", "BLOCKING") for f in report.findings):
        raise typer.Exit(1)

@app.command()
def reproduce(file_path: Path):
    """Attempt to cleanly reproduce a calculation (specifically notebooks)."""
    if not file_path.exists():
        console.print(f"[red]File {file_path} not found.[/red]")
        raise typer.Exit(1)
        
    if file_path.suffix == ".ipynb":
        console.print(f"Executing notebook {file_path} from a clean kernel...")
        auditor = NotebookAuditor(file_path)
        report = auditor.audit(reproduce=True)
        print_report(report)
    else:
        console.print("[yellow]Reproduce is currently only supported for Jupyter Notebooks.[/yellow]")

@app.command()
def demo():
    """Run a flagship demo of the auditor."""
    console.print(Panel("[bold magenta]scientific-calculation-auditor[/bold magenta]\nReconstructing the chain of evidence."))
    
    import tempfile
    
    console.print("\n[bold cyan]--- Python Auditor Demo ---[/bold cyan]")
    code = (
        "# calculation.py\n"
        "import math\n\n"
        "g = 9.81 # undocumented constant\n"
        "flow_rate_m3_h = 2500\n"
        "flow_rate_m3_s = flow_rate_m3_h / 3600 # inline conversion\n\n"
        "# early rounding\n"
        "pressure_drop = round(0.5 * 1000 * flow_rate_m3_s**2, 2)\n"
        "print(pressure_drop)\n"
    )
    
    with tempfile.NamedTemporaryFile("w", suffix=".py", delete=False) as f:
        f.write(code)
        path = Path(f.name)
        
    console.print("Auditing a thermal/hydraulic Python script with deliberate errors:")
    console.print(f"[dim]{code}[/dim]")
    
    report = PythonAuditor(path).audit()
    print_report(report)
    path.unlink()
    
    console.print("\n[bold cyan]--- Spreadsheet Auditor Demo ---[/bold cyan]")
    # Generate a dummy xlsx
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sizing"
    ws["A1"] = "=B1*C1"
    ws["A2"] = "=B2*C2"
    ws["A3"] = 42.5  # Manual override
    ws["A4"] = "=B4*C4"
    ws["B5"] = "=A1*9.81" # Hardcoded constant
    ws["C6"] = "=[External.xlsx]Sheet1!A1" # External link
    
    with tempfile.NamedTemporaryFile("wb", suffix=".xlsx", delete=False) as f:
        wb.save(f.name)
        path_xlsx = Path(f.name)
        
    console.print("Auditing an Excel spreadsheet with deliberate errors...")
    report2 = SpreadsheetAuditor(path_xlsx).audit()
    print_report(report2)
    path_xlsx.unlink()

def main():
    app()

if __name__ == "__main__":
    main()
